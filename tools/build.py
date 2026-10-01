"""Assemble the full manuscript: approved Ch1-3 (copied verbatim from the client's docx)
+ revised Ch4-18 (manuscript/chXX.md), styled to match the approved chapters."""
import copy, re, sys, glob, os
import docx
from docx.enum.text import WD_BREAK

SRC = sys.argv[1]          # approved Revised Ch1-3 docx
OUT_DOCX = sys.argv[2]
OUT_MD = sys.argv[3]
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

d = docx.Document(SRC)
P = d.paragraphs
# templates (deep copies of approved formatting)
T = {
    'heading': copy.deepcopy(P[111]._p),   # "Chapter Two"
    'first':   copy.deepcopy(P[5]._p),     # first para, no indent
    'body':    copy.deepcopy(P[6]._p),     # indented para
    'break':   copy.deepcopy(P[9]._p),     # * * *
    'loghead': copy.deepcopy(P[97]._p),    # #camille-case header
    'logline': copy.deepcopy(P[98]._p),    # Wren, 7:41 AM: ...
    'spacer':  copy.deepcopy(P[107]._p),   # empty after log
}
body = d.element.body
sectPr = body[-1] if body[-1].tag.endswith('sectPr') else None

# front matter: series name, drop the "Revised manuscript" line
P[1].runs[0].text = 'Quiet Hollow Mystery — Book One'
for r in P[1].runs[1:]:
    pass
P[3]._p.getparent().remove(P[3]._p)

def new_para(kind):
    el = copy.deepcopy(T[kind])
    # strip runs, keep pPr
    for child in list(el):
        if not child.tag.endswith('}pPr'):
            el.remove(child)
    if sectPr is not None:
        sectPr.addprevious(el)
    else:
        body.append(el)
    return docx.text.paragraph.Paragraph(el, d._body)

def add_runs(par, text, mono=False, bold_prefix=None, all_bold=False):
    parts = []
    if bold_prefix:
        parts.append((bold_prefix, True, False))
        text = text[len(bold_prefix):]
    # italics via *...*
    for i, chunk in enumerate(re.split(r'\*([^*]+)\*', text)):
        if chunk:
            parts.append((chunk, all_bold, i % 2 == 1))
    for t, b, it in parts:
        r = par.add_run(t)
        if b: r.bold = True
        if it: r.italic = True
        if mono:
            r.font.name = 'Courier New'
            r.font.size = docx.shared.Pt(10)

USER = re.compile(r'^((?:Wren|Des), \d{1,2}:\d{2} [AP]M(?: \([^)]*\))?:|(?:me|Des|Priya|Theo|Reyes|Kenzie):|[a-z][a-z0-9._]{2,}:)')

def log_line(line):
    if line.startswith('#'):
        p = new_para('loghead'); add_runs(p, line, mono=True, all_bold=True); return
    p = new_para('logline')
    m = USER.match(line)
    caps = line.strip() and line.strip() == line.strip().upper() and any(c.isalpha() for c in line) and len(line) < 70 and not line.startswith('-')
    if m:
        add_runs(p, line, mono=True, bold_prefix=m.group(1))
    else:
        add_runs(p, line, mono=True, all_bold=bool(caps))

def page_break_para():
    p = new_para('first'); p.add_run().add_break(WD_BREAK.PAGE)

md_out = []
chapters = sorted(glob.glob(os.path.join(ROOT, 'manuscript', 'ch*.md')))
for path in chapters:
    lines = open(path, encoding='utf-8').read().split('\n')
    prev = 'heading'
    in_log = False
    buf = []
    for raw in lines:
        line = raw.rstrip()
        md_out.append(raw)
        if line == '[[LOG]]':
            in_log = True; continue
        if line == '[[/LOG]]':
            in_log = False; new_para('spacer'); prev = 'log'; continue
        if in_log:
            if line.strip(): log_line(line)
            continue
        if not line.strip():
            continue
        if line.startswith('# '):
            p = new_para('heading'); p.add_run(line[2:]); prev = 'heading'; continue
        if line.strip() == '* * *':
            p = new_para('break'); p.add_run('*   *   *'); prev = 'break'; continue
        kind = 'first' if prev in ('heading', 'break', 'log') else 'body'
        p = new_para(kind); add_runs(p, line); prev = 'para'
    md_out.append('')

# Book Two preview placeholder
p = new_para('heading'); p.add_run('BOOK TWO PREVIEW')
p = new_para('first'); p.alignment = 1; r = p.add_run('Dead Reckoning'); r.italic = True; r.font.size = docx.shared.Pt(16)
p = new_para('first'); p.alignment = 1; p.add_run('Chapter One')
p = new_para('first'); p.alignment = 1; r = p.add_run('[Placeholder: the first chapter of Dead Reckoning will be supplied and inserted here.]'); r.italic = True
md_out += ['# BOOK TWO PREVIEW', '', '*Dead Reckoning*', '', 'Chapter One', '', '*[Placeholder: the first chapter of Dead Reckoning will be supplied and inserted here.]*', '']

d.core_properties.title = 'The Wrong Kind of Missing'
d.core_properties.author = 'J. Lowery'
d.save(OUT_DOCX)

# Markdown edition: Ch1-3 text from the approved docx + Ch4-18
src = docx.Document(SRC)
head = ['# THE WRONG KIND OF MISSING', '', '*Quiet Hollow Mystery — Book One*', '', 'J. Lowery', '']
for para in src.paragraphs[4:]:
    t = ''.join(('*%s*' % r.text if r.italic and r.text.strip() else r.text) for r in para.runs)
    if para.style is not None and para.style.name.startswith('Heading'):
        head += ['# ' + t, '']
    elif t.strip() == '*   *   *':
        head += ['* * *', '']
    elif t.strip():
        head += [t, '']
open(OUT_MD, 'w', encoding='utf-8').write('\n'.join(head + md_out))
print('ok', len(chapters), 'chapters appended')
